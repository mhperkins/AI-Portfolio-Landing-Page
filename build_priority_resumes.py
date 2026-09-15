"""Build the three priority application resumes from Priority Applications Resumes.md.

Edit the markdown, then run this. Each '# RESUME: <name>' block becomes
'Maxwell Perkins Resume - <name>.docx' in the academic CV's docx format.
"""
import copy
import re
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

ROOT = r"C:\Users\maxwe\OneDrive\Desktop\Claude\Apps and Tools\AI Tool Portfolio"
SOURCE = ROOT + r"\Priority Applications Resumes.md"
# Word locks a docx while it is open; the Community Engagement resume carries the same styles.
TEMPLATES = [ROOT + r"\Academic CV - UW Madison.docx",
             ROOT + r"\Maxwell Perkins Resume - Community Engagement.docx"]

LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_LEAD = re.compile(r"^\*\*(.+?)\*\*\s*(.*)$")


def load_template():
    for path in TEMPLATES:
        try:
            return Document(path)
        except PermissionError:
            continue
    sys.exit("Both templates are locked. Close them in Word and rerun.")


def build(name, lines):
    doc = load_template()
    border = None
    for p in doc.paragraphs:
        if p.text.strip() == "EDUCATION":
            border = copy.deepcopy(p._p.pPr.find(qn("w:pBdr")))
            break
    doc.styles["Normal"].font.size = Pt(10.5)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Pt(46)
    sec.top_margin = sec.bottom_margin = Pt(36)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    def para(style="Normal", align=None, before=0, after=2):
        p = doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        return p

    def run(p, text, bold=False, italic=False, size=None):
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        return r

    def link(p, text, url):
        rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = p.add_run(text)
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        h.append(r._r)
        p._p.append(h)

    last = None  # kind of the previous element, for spacing
    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("Contact: "):
            run(para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2), "Maxwell Perkins", bold=True, size=20)
            text = line[len("Contact: "):].replace(" | ", "  |  ")
            run(para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0), text)
            last = "header"
        elif line.startswith("Links: "):
            p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
            for i, (text, url) in enumerate(LINK.findall(line)):
                if i:
                    run(p, "  |  ")
                link(p, text, url)
            last = "header"
        elif line.startswith("## "):
            p = para(before=8, after=3)
            run(p, line[3:], bold=True, size=12)
            if border is not None:
                p._p.get_or_add_pPr().append(copy.deepcopy(border))
            last = "heading"
        elif line.startswith("- "):
            p = para(style="List Bullet", after=1)
            m = BOLD_LEAD.match(line[2:])
            if m:
                run(p, m.group(1) + " ", bold=True)
                run(p, m.group(2))
            else:
                run(p, line[2:])
            last = "bullet"
        elif line.startswith("**") and "** |" in line:
            title, meta = line[2:].split("** |", 1)
            before = {"heading": 2, "entry": 0}.get(last, 4)
            p = para(before=before, after=1)
            run(p, title, bold=True)
            run(p, "  |  " + "  |  ".join(part.strip() for part in meta.split("|")))
            last = "entry"
        elif line.startswith("*") and not line.startswith("**"):
            run(para(after=1), line.strip("*"), italic=True)
            last = "note"
        elif BOLD_LEAD.match(line):
            m = BOLD_LEAD.match(line)
            p = para(before=2 if last == "heading" else 0, after=1)
            run(p, m.group(1) + " ", bold=True)
            run(p, m.group(2))
            last = "line"
        else:
            sys.exit(f"{name}: unrecognized line: {line}")

    out = ROOT + rf"\Maxwell Perkins Resume - {name}.docx"
    doc.save(out)
    print("saved", out)


text = open(SOURCE, encoding="utf-8").read().split("\n# NOTES", 1)[0]
parts = re.split(r"^# RESUME: (.+)$", text, flags=re.M)
for i in range(1, len(parts), 2):
    build(parts[i].strip(), parts[i + 1].splitlines())
