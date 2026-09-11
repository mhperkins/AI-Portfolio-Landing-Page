"""Build the community engagement resume in the academic CV's docx format."""
import copy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

ROOT = r"C:\Users\maxwe\OneDrive\Desktop\Claude\Apps and Tools\AI Tool Portfolio"
TEMPLATE = ROOT + r"\Academic CV - UW Madison.docx"
OUT = ROOT + r"\Maxwell Perkins Resume - Community Engagement.docx"

doc = Document(TEMPLATE)

# Lift the section-heading border from the template before clearing the body.
heading_border = None
for p in doc.paragraphs:
    if p.text.strip() == "EDUCATION":
        heading_border = copy.deepcopy(p._p.pPr.find(qn("w:pBdr")))
        break

doc.styles["Normal"].font.size = Pt(10.5)
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Pt(46)
sec.top_margin = sec.bottom_margin = Pt(36)

body = doc.element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)


def para(style="Normal", align=None, before=0, after=2):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def run(p, text, bold=False, italic=False, size=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return r


def link(p, text, url, size=None):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    if size:
        r.font.size = Pt(size)
    h.append(r._r)
    p._p.append(h)


def heading(text):
    p = para(before=8, after=3)
    run(p, text, bold=True, size=12)
    if heading_border is not None:
        p._p.get_or_add_pPr().append(copy.deepcopy(heading_border))


def entry(title, meta, before=4):
    p = para(before=before, after=1)
    run(p, title, bold=True)
    run(p, "  |  " + meta)


def note(text):
    p = para(after=1)
    run(p, text, italic=True)


def bullet(text, lead=None):
    p = para(style="List Bullet", after=1)
    if lead:
        run(p, lead, bold=True)
    run(p, text)


# Header
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
run(p, "Maxwell Perkins", bold=True, size=20)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
run(p, "Brooklyn, NY (relocating to Madison, WI)  |  maxwellhperkins@gmail.com  |  925-259-1748", size=10.5)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
link(p, "linkedin.com/in/maxwell-perkins-702a82207",
     "https://www.linkedin.com/in/maxwell-perkins-702a82207/", size=10.5)

heading("EDUCATION")
entry("M.M., Music Theory and Composition (Songwriting Concentration)", "New York University, 2022", before=2)
entry("B.A., Music Composition, Minor in Business", "University of Wisconsin-Madison, 2014", before=0)

heading("TEACHING EXPERIENCE")
entry("Private Music Instructor", "2011-Present", before=2)
bullet("Teach piano, drums, audio production, and composition to students ages 5 to 18, one-on-one and in groups.")
bullet("Schedule lessons and keep families informed on each student's progress.")
bullet("Developed original curricula for composition and audio production, from beginner through advanced.")

heading("PROFESSIONAL EXPERIENCE")
entry("Community Manager, Sprout Society", "Brooklyn, NY  |  February 2026-Present", before=2)
note("Nonprofit community space for creatives. Team of three.")
bullet("Serve as the first point of contact for members, partners, and donors, responding to every inquiry.")
bullet("Maintain the organization's CRM and records, keeping member, donor, partner, and event data accurate and current.")
bullet("Built and run the sign-up and onboarding process for new members.")
bullet("Plan and host a monthly music and art showcase, and support members running their own events in the space.")
bullet("Grew the space from one showcase a month to 10 to 15 events, bringing in over 100 new people each month.")
bullet("Lead social media and community outreach; support grant research, applications, and reporting.")

entry("Founder, Music Major Records LLC", "Jul 2023-Aug 2025")
bullet("Ran and iterated a platform for college musicians to build careers, then placed it on hold.")

entry("Producer & Recording Engineer, Misfits' Instruments Studio", "2018-Present")
bullet("Produce, arrange, and mix recordings from tracking through final mix in Pro Tools and Logic.")
bullet("Handle studio scheduling and session logistics, and play as a session musician.")

entry("Composer and Songwriter, Independent", "2014-Present")
bullet("Write pop and concert music, with songwriting and producer credits for John Harvie and Brett Altman.")
bullet("Produce and host a concert music series for composers and performers in NYC.")

heading("APPLIED SYSTEMS")
entry("The Sprout Suite", "2025-Present", before=2)
note("Tools I built to keep Sprout Society's records accurate, cut manual data entry, and replace paid software.")
bullet("manages contacts, members, donors, outreach, and events in one place.", lead="CRM: ")
bullet("grant research and matching, application tracking through the full grant cycle, and a writing workspace.",
       lead="Grant Tool: ")
bullet("a content library and workspace that publishes posts automatically on a content calendar.",
       lead="Social Media Manager: ")
bullet("builds dynamic QR codes, tracks scans, and visualizes the results.", lead="Campaign Tracker: ")

heading("ADDITIONAL")
p = para(before=2, after=1)
run(p, "Tools: ", bold=True)
run(p, "Google Workspace, CRM systems, Canva, social media (Instagram, YouTube, TikTok)")
p = para(after=0)
run(p, "Availability: ", bold=True)
run(p, "Days, plus nights and weekends when needed; valid driver's license")

doc.save(OUT)
print("saved", OUT)
