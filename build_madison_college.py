"""Build the Madison College CV and cover letter in the academic CV's docx format."""
import copy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

ROOT = r"C:\Users\maxwe\OneDrive\Desktop\Claude\Apps and Tools\AI Tool Portfolio"
TEMPLATE = ROOT + r"\Academic CV - UW Madison.docx"
CV_OUT = ROOT + r"\Maxwell Perkins CV - Madison College.docx"
LETTER_OUT = ROOT + r"\Maxwell Perkins Cover Letter - Madison College.docx"

CONTACT = "Brooklyn, NY (relocating to Madison, WI)  |  maxwellhperkins@gmail.com  |  925-259-1748"


def fresh(size, margin_lr, margin_tb):
    doc = Document(TEMPLATE)
    border = None
    for p in doc.paragraphs:
        if p.text.strip() == "EDUCATION":
            border = copy.deepcopy(p._p.pPr.find(qn("w:pBdr")))
            break
    doc.styles["Normal"].font.size = Pt(size)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Pt(margin_lr)
    sec.top_margin = sec.bottom_margin = Pt(margin_tb)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return doc, border


def para(doc, style="Normal", align=None, before=0, after=2):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def run(p, text, bold=False, italic=False, size=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return r


def link(p, text, url, size=None):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    if size:
        r.font.size = Pt(size)
    h.append(r._r)
    p._p.append(h)


def letterhead(doc, size):
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    run(p, "Maxwell Perkins", bold=True, size=20)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    run(p, CONTACT, size=size)
    p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    link(p, "linkedin.com/in/maxwell-perkins-702a82207",
         "https://www.linkedin.com/in/maxwell-perkins-702a82207/", size=size)


# ---------- CV ----------
doc, border = fresh(10.5, 46, 36)


def heading(text):
    p = para(doc, before=8, after=3)
    run(p, text, bold=True, size=12)
    if border is not None:
        p._p.get_or_add_pPr().append(copy.deepcopy(border))


def entry(title, meta, before=4):
    p = para(doc, before=before, after=1)
    run(p, title, bold=True)
    run(p, "  |  " + meta)


def note(text):
    p = para(doc, after=1)
    run(p, text, italic=True)


def bullet(text, lead=None):
    p = para(doc, style="List Bullet", after=1)
    if lead:
        run(p, lead, bold=True)
    run(p, text)


def line(label, text, after=1):
    p = para(doc, after=after)
    run(p, label, bold=True)
    run(p, text)


letterhead(doc, 10.5)

heading("EDUCATION")
entry("M.M., Music Theory and Composition (Songwriting Concentration)", "New York University, 2022", before=2)
entry("B.A., Music Composition, Minor in Business", "University of Wisconsin-Madison, 2014", before=0)

heading("TEACHING EXPERIENCE")
entry("Private Music Instructor", "2011-Present", before=2)
bullet("Teach piano, drums, audio production, and composition to students ages 5 to 18, one-on-one and in groups.")
bullet("Plan mixed-level group lessons so beginning and advanced students learn in the same session.")
bullet("Developed original curricula for composition and audio production, from beginner through advanced.")
bullet("Schedule lessons and keep families informed on each student's progress.")

heading("MUSIC EXPERIENCE")
entry("Composer and Songwriter, Independent", "2014-Present", before=2)
bullet("Write pop, rock, and concert music, with songwriting and producer credits for John Harvie and Brett Altman.")
bullet("Returned to concert music with a re-debut concert in NYC in June 2026.")
bullet("Produce and host a concert music series for composers and performers in NYC.")

entry("Producer & Recording Engineer, Misfits' Instruments Studio", "2018-Present")
bullet("Produce, arrange, and mix recordings from tracking through final mix in Pro Tools and Logic.")
bullet("Play as a session musician, and handle studio scheduling and session logistics.")

heading("PROFESSIONAL EXPERIENCE")
entry("Community Manager, Sprout Society", "Brooklyn, NY  |  February 2026-Present", before=2)
note("Nonprofit community space for creatives. Team of three.")
bullet("Plan and host a monthly music and art showcase, and support members running their own events in the space.")
bullet("Grew the space from one showcase a month to 10 to 15 events, bringing in over 100 new people each month.")
bullet("Serve as the first point of contact for members, partners, and donors; built the sign-up and onboarding process.")

entry("Freelance Prompt Engineer, Handshake AI", "Aug 2025-Present")
bullet("Write rubrics to evaluate musical analysis, identify errors, and document the correct answer.")

entry("Founder, Music Major Records LLC", "Jul 2023-Aug 2025")
bullet("Built and ran a platform for college musicians to collaborate and build their careers.")

heading("APPLIED SYSTEMS")
note("Software I designed and built.")
bullet("a studio manager for music teachers: scheduling, student records, booking, and pay tracking.",
       lead="TeacherAID: ")
bullet("a composition tool that gives feedback on each draft, from first idea to finished score.",
       lead="The Composer Compass: ")
bullet("a CRM, grant tool, social media manager, and campaign tracker for Sprout Society.",
       lead="The Sprout Suite: ")

heading("ADDITIONAL")
p = para(doc, before=2, after=1)
run(p, "Instruments and software: ", bold=True)
run(p, "Piano, drums, Pro Tools, Logic, Google Workspace, Canva")
line("Availability: ", "Day, evening, and weekend sections; valid driver's license", after=0)

doc.save(CV_OUT)
print("saved", CV_OUT)

# ---------- Cover letter ----------
LETTER = [
    "After nearly ten years in New York, I'm moving back to Madison, and I'd love to teach in Madison "
    "College's Music program. I earned my B.A. in composition at UW-Madison and my M.M. in Music Theory "
    "and Composition at NYU, and I've taught music for fifteen years.",

    "I teach piano, drums, audio production, and composition to students ages 5 to 18, one-on-one and in "
    "groups. Group lessons taught me the core skill of a mixed-level classroom: one plan that reaches a "
    "beginner and an advanced student in the same session. I also wrote my own curricula for composition "
    "and audio production. Madison College would be my first college classroom, and I'd bring the same "
    "approach I use in every lesson: find exactly where a student's understanding breaks down, and build "
    "from there.",

    "I can contribute most in Music Theory Fundamentals and History of Pop and Rock Music. My master's "
    "concentrated in songwriting, and for more than a decade I've written, produced, and recorded pop and "
    "rock as a songwriter, session musician, and engineer. I also write concert music and host a concert "
    "series for composers and performers in New York, which gives me a broad base for Music Appreciation. "
    "Pop and rock history starts with the blues, gospel, and R&B, and I'd make sure students hear those "
    "Black American roots directly, not as background.",

    "As Community Manager at Sprout Society, a nonprofit community space for creatives in Brooklyn, I plan "
    "and host a monthly music and art showcase. We've grown from one event a month to 10 to 15, bringing in "
    "over 100 new people each month. That work comes down to making every newcomer feel they belong in the "
    "room, whatever their background or experience, and I want the same for every student in my sections.",

    "I'm also comfortable with educational technology. I design and build software for teaching and "
    "composing, including a studio manager for independent music teachers and a composition tool that "
    "guides a writer from first idea to finished score. I'd bring that same care to clear, well-organized "
    "materials for face-to-face, hybrid, and online sections.",

    "I'm available for day, evening, and weekend sections, and I look forward to discussing how I can "
    "support your students and the Music program.",
]

doc, _ = fresh(11, 72, 54)
letterhead(doc, 10.5)
para(doc, before=14, after=10).add_run("September 14, 2026")
para(doc, after=10).add_run("Dear Hiring Committee,")
for text in LETTER:
    para(doc, after=10).add_run(text)
para(doc, before=4, after=18).add_run("Thank you for your time and consideration,")
run(para(doc, after=0), "Maxwell Perkins", bold=True)
para(doc, after=0).add_run("maxwellhperkins@gmail.com  |  925-259-1748")

doc.save(LETTER_OUT)
print("saved", LETTER_OUT)
