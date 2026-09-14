"""Build the AI roles CV in the academic CV's docx format."""
import copy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt, RGBColor

ROOT = r"C:\Users\maxwe\OneDrive\Desktop\Claude\Apps and Tools\AI Tool Portfolio"
TEMPLATE = ROOT + r"\Academic CV - UW Madison.docx"
# Word locks the CV while it is open; the community engagement resume carries the same styles.
FALLBACK = ROOT + r"\Maxwell Perkins Resume - Community Engagement.docx"
OUT = ROOT + r"\Maxwell Perkins CV - AI Roles.docx"

try:
    doc = Document(TEMPLATE)
except Exception:
    print("template locked, using", FALLBACK)
    doc = Document(FALLBACK)

# Lift the section-heading border from the template before clearing the body.
heading_border = None
for p in doc.paragraphs:
    if p.text.strip() == "EDUCATION":
        heading_border = copy.deepcopy(p._p.pPr.find(qn("w:pBdr")))
        break

doc.styles["Normal"].font.size = Pt(10)
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Pt(42)
sec.top_margin = sec.bottom_margin = Pt(26)

body = doc.element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)


def para(style="Normal", align=None, before=0, after=2):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    return p


def run(p, text, bold=False, italic=False, size=None):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    return r


def link(p, text, url, size=None):
    rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    if size:
        r.font.size = Pt(size)
    h.append(r._r)
    p._p.append(h)


def heading(text):
    p = para(before=5, after=3)
    run(p, text, bold=True, size=11.5)
    if heading_border is not None:
        p._p.get_or_add_pPr().append(copy.deepcopy(heading_border))


def entry(title, meta, before=4):
    p = para(before=before, after=1)
    run(p, title, bold=True)
    run(p, "  |  " + meta)


def note(text, before=0):
    p = para(before=before, after=1)
    run(p, text, italic=True)


def bullet(text, lead=None):
    p = para(style="List Bullet", after=1)
    if lead:
        run(p, lead, bold=True)
    run(p, text)


def line(lead, text, before=0, after=1):
    p = para(before=before, after=after)
    run(p, lead, bold=True)
    run(p, text)


# Header
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
run(p, "Maxwell Perkins", bold=True, size=20)
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
run(p, "Brooklyn, NY (relocating to Madison, WI)  |  maxwellhperkins@gmail.com  |  925-259-1748")
p = para(align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
link(p, "Portfolio", "https://ai-portfolio-landing-page.vercel.app/")
run(p, "  |  ")
link(p, "LinkedIn", "https://www.linkedin.com/in/maxwell-perkins-702a82207/")
run(p, "  |  ")
link(p, "GitHub", "https://github.com/mhperkins")

heading("SUMMARY")
p = para(before=2, after=1)
run(p, "AI tool and workflow builder. Ships production internal tools on Next.js and Supabase, Claude API "
       "and MCP integrations, and structured prompts that return data an app can import directly. "
       "Background in LLM evaluation, nonprofit operations, and music.")

heading("SKILLS")
line("AI: ", "Claude API, Claude Code, MCP servers, prompt and rubric design, LLM evaluation, structured JSON "
             "output, RAG (ChromaDB)", before=2)
line("Automation: ", "n8n, scheduled jobs (pg_cron, Vercel cron), REST APIs, Instagram API (via Composio), "
                     "Python scripting")
line("Web and data: ", "Next.js, React, JavaScript, TypeScript, Tailwind CSS, Supabase (PostgreSQL, Auth, RLS), "
                       "SQL, Vercel, Git", after=0)

heading("APPLIED SYSTEMS")
entry("The Sprout Suite", "2025-Present", before=2)
note("Four production tools for a Brooklyn nonprofit on one Supabase backend, replacing paid software.")
bullet("exposes the live CRM to Claude as 14 tools (relationship health, search, touchpoint logging, dedupe).",
       lead="MCP server: ")
bullet("multi-phase prompts that return schema-validated JSON the CRM and Grant Tool import directly.",
       lead="Research protocols: ")
bullet("posts move from draft to approved to scheduled, then publish to Instagram with human sign-off.",
       lead="Social Media Manager: ")
bullet("tracks each grant from research brief to submission, with versioned answer drafts, tasks, and export.",
       lead="Grant Tool: ")
bullet("QR codes that can be retargeted after printing, with scan counts on Vercel and Supabase.",
       lead="Campaign Tracker: ")
bullet("built an n8n workflow that pulled data from all four tools and added a Claude analysis layer.",
       lead="Weekly digest: ")

entry("TeacherAID", "2026-Present")
note("Scheduling, payroll, and booking for an independent music teacher, in production on Next.js, Supabase, "
     "and Vercel.")
bullet("Logs each lesson in one tap; hours, pay sheets, and student records all update from that entry.")
bullet("Sends scheduled email summaries through pg_cron and secured API routes, with Claude observations on "
       "each pay period.")
bullet("Lets families book open lesson times from a single link.")

heading("AI PROJECTS: THE COMPOSER HUB")
note("AI tools that assist trained musicians in their process instead of generating music for them.", before=2)
entry("The Composer Compass", "2026-Present", before=2)
bullet("Evaluates each piece with an analysis agent, a structured intake, and a rubric of six dimensions plus "
       "three modifiers.")
bullet("Grounds evaluations in the composer's theory library (ChromaDB retrieval) and compares models on one "
       "weighted rubric.")
bullet("Maps the workflow from planning through sketching, drafts, and engraving in a Next.js canvas built on "
       "the Anthropic API.")
entry("Composer CRM", "2026-Present")
bullet("Tracks performers, ensembles, venues, and concerts, with a newsletter builder and AI prospect research.")
entry("Virtual Agency", "2026-Present")
bullet("Claude Code agents, each a system prompt plus tools, that have produced a landing page, program notes, "
       "and a campaign.")

heading("PROFESSIONAL EXPERIENCE")
entry("Community Manager, Sprout Society", "Brooklyn, NY  |  February 2026-Present", before=2)
bullet("Run operations for a three-person nonprofit: membership, 10 to 15 events a month, outreach, social "
       "media, and grants.")
bullet("Scoped, built, and maintain the organization's internal tools (see Applied Systems).")
entry("Freelance Prompt Engineer, Handshake AI", "Aug 2025-Present")
bullet("Annotate AI-generated music; write prompts and rubrics that expose failures in AI musical analysis.")
entry("Founder, Music Major Records LLC", "Jul 2023-Aug 2025")
bullet("Owned strategy, product, marketing, and finances for a platform for college musicians, then placed it "
       "on hold.")
entry("Private Music Instructor", "2011-Present")
bullet("Piano, drums, audio production, and composition, one-on-one and in groups, with original curricula.")
entry("Producer & Recording Engineer, Misfits' Instruments Studio", "2018-Present")

heading("EDUCATION")
entry("M.M., Music Theory and Composition (Songwriting Concentration)", "New York University, 2022", before=2)
entry("B.A., Music Composition, Minor in Business", "University of Wisconsin-Madison, 2014", before=0)

doc.save(OUT)
print("saved", OUT)
